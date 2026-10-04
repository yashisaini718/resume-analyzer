import io
import docx
import pytest
from unittest.mock import MagicMock, patch
from fastapi.testclient import TestClient

from backend.src.services.resume_parser import clean_resume_text, parse_and_clean_resume_file
from backend.src.ai.prompts import build_analysis_prompt
from backend.src.schemas.analyzer import (
    ResumeAnalysisResult,
    MissingRequirement,
    InterviewQuestion,
    RoadmapStep,
)
from backend.src.main import app

client = TestClient(app)


def test_clean_resume_text():
    dirty_text = "• Senior  Python Engineer\n\n\n\n• Experience with   FastAPI &   Docker\u2022 AWS\u2013Cloud"
    cleaned = clean_resume_text(dirty_text)
    assert "- Senior Python Engineer" in cleaned
    assert "- Experience with FastAPI & Docker" in cleaned
    assert "- AWS-Cloud" in cleaned
    assert "\n\n\n" not in cleaned


def test_parse_docx_file():
    doc = docx.Document()
    doc.add_heading("John Doe - Software Engineer", level=1)
    doc.add_paragraph("Proficient in Python, Go, and Kubernetes.")
    stream = io.BytesIO()
    doc.save(stream)
    docx_bytes = stream.getvalue()

    cleaned = parse_and_clean_resume_file("resume.docx", docx_bytes)
    assert "John Doe" in cleaned
    assert "Proficient in Python, Go, and Kubernetes." in cleaned


def test_parse_txt_file():
    txt_bytes = "Jane Doe\nBackend Developer with 4 years experience in Django and PostgreSQL.".encode("utf-8")
    cleaned = parse_and_clean_resume_file("resume.txt", txt_bytes)
    assert "Backend Developer with 4 years experience" in cleaned


def test_build_analysis_prompt():
    prompt = build_analysis_prompt(
        resume_text="Resume: Python Developer",
        job_description="JD: Senior Backend Lead with Kafka & Golang",
        self_description="Self: Passionate about distributed systems",
    )
    assert "TARGET JOB DESCRIPTION" in prompt
    assert "CANDIDATE SELF-DESCRIPTION" in prompt
    assert "CLEANED RESUME CONTENT" in prompt
    assert "Kafka & Golang" in prompt


def test_schema_validation():
    mock_data = {
        "profile_summary": "Strong candidate with solid Python background but lacking Go & Kafka experience.",
        "overall_match_score": 75,
        "missing_requirements": [
            {
                "requirement": "Hands-on Go / Golang microservices",
                "category": "Technical Skill",
                "importance": "Critical",
                "gap_reasoning": "Job description requires 3+ years Go; candidate has exclusively Python background.",
                "fix_recommendation": "Build a Go-based distributed message consumer project.",
            }
        ],
        "interview_questions": [
            {
                "question_number": i,
                "question": f"Question {i} regarding distributed architectures?",
                "question_type": "Technical Deep-Dive",
                "why_asked": "Assesses knowledge of messaging patterns.",
                "key_points_to_cover": ["Eventual consistency", "Idempotency", "DLQ patterns"],
            }
            for i in range(1, 6)
        ],
        "preparation_roadmap": [
            {
                "phase_number": 1,
                "phase_name": "Go Syntax & Concurrency",
                "timeline": "Week 1-2",
                "objectives": ["Master Goroutines and channels"],
                "action_items": ["Complete Tour of Go", "Implement worker pool pattern"],
                "recommended_projects_or_resources": ["Go in Action", "Build high-throughput API"],
            }
        ],
    }

    result = ResumeAnalysisResult.model_validate(mock_data)
    assert result.overall_match_score == 75
    assert len(result.interview_questions) == 5
    assert len(result.missing_requirements) == 1


def test_health_check_endpoint():
    response = client.get("/health")
    assert response.status_code == 200
    assert response.json()["status"] == "healthy"


def test_parse_file_endpoint():
    txt_content = b"Resume: Data Engineer skilled in Apache Spark, PySpark, and AWS Redshift."
    response = client.post(
        "/api/v1/analyzer/parse-file",
        files={"resume_file": ("resume.txt", io.BytesIO(txt_content), "text/plain")},
    )
    assert response.status_code == 200
    data = response.json()
    assert data["success"] is True
    assert "Apache Spark" in data["cleaned_text"]


def test_analyze_text_mocked():
    sample_result = ResumeAnalysisResult(
        profile_summary="Candidate is a mid-level Python developer applying for a Senior AI Engineer role.",
        overall_match_score=68,
        missing_requirements=[
            MissingRequirement(
                requirement="PyTorch model fine-tuning",
                category="Technical Skill",
                importance="Critical",
                gap_reasoning="Resume lacks production PyTorch training experience.",
                fix_recommendation="Fine-tune a LoRA model on HuggingFace and document in portfolio.",
            )
        ],
        interview_questions=[
            InterviewQuestion(
                question_number=i,
                question=f"Test Question {i} on LLM fine-tuning?",
                question_type="Technical Deep-Dive",
                why_asked="Probing model training depth.",
                key_points_to_cover=["LoRA rank", "learning rate schedules", "loss evaluation"],
            )
            for i in range(1, 6)
        ],
        preparation_roadmap=[
            RoadmapStep(
                phase_number=1,
                phase_name="LLM Fine-tuning Mastery",
                timeline="Week 1-2",
                objectives=["Understand PEFT and QLoRA"],
                action_items=["Fine-tune open weights model"],
                recommended_projects_or_resources=["HuggingFace PEFT docs"],
            )
        ],
    )

    with patch("backend.src.ai.pipeline.pipeline.analyze", return_value=sample_result):
        payload = {
            "resume_text": "Experienced Python developer with 3 years building web apps using Django and FastAPI.",
            "job_description": "Senior AI Engineer required with deep PyTorch, LLM fine-tuning, and LangChain knowledge.",
            "self_description": "I have built standard apps and have started learning generative AI.",
        }
        response = client.post("/api/v1/analyzer/analyze-text", json=payload)
        assert response.status_code == 200
        res_json = response.json()
        assert res_json["success"] is True
        assert res_json["data"]["overall_match_score"] == 68
        assert len(res_json["data"]["missing_requirements"]) == 1
        assert len(res_json["data"]["interview_questions"]) == 5
        assert len(res_json["data"]["preparation_roadmap"]) == 1
